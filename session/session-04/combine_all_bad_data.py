import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = r"d:\GoogleDrive\VibeCoding\teacher\session\session-04"

# Import helper functions and data generators from generate_bad_data_files
import sys
sys.path.append(OUTPUT_DIR)
import generate_bad_data_files as src

def style_index_sheet(ws):
    ws.views.sheetView[0].showGridLines = True
    
    title_font = Font(name="Noto Sans Thai", size=15, bold=True, color="1E1B4B")
    sub_font = Font(name="Noto Sans Thai", size=10, italic=True, color="4338CA")
    header_font = Font(name="Noto Sans Thai", size=10, bold=True, color="FFFFFF")
    cell_font = Font(name="Noto Sans Thai", size=10, color="1E293B")
    bold_cell_font = Font(name="Noto Sans Thai", size=10, bold=True, color="0F172A")
    tag_font = Font(name="Noto Sans Thai", size=9, bold=True, color="1D4ED8")
    
    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    tag_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    ws["B2"] = "📚 รวมชุดข้อมูลตัวอย่าง 7 กับดัก Messy Data (Session 04 Master Collection)"
    ws["B2"].font = title_font
    
    ws["B3"] = "รายวิชา: Fundamentals of Data Science | คณะการสร้างเจ้าของธุรกิจ มหาวิทยาลัยศรีปทุม"
    ws["B3"].font = sub_font
    
    ws["B4"] = "ชุดข้อมูลจำลองปัญหาการเก็บข้อมูลในธุรกิจไทย 7 ประเภท เพื่อฝึกฝนการจับผิดตารางเละและแปลงสู่โครงสร้าง Tidy Data"
    ws["B4"].font = Font(name="Noto Sans Thai", size=9.5, color="475569")
    
    headers = [
        "ลำดับ",
        "กับดักข้อผิดพลาด (Trap Name)",
        "กรณีศึกษาธุรกิจไทย",
        "แผ่นงานข้อมูลดิบ (Data Sheet)",
        "แผ่นงานคำอธิบาย (Explanation Sheet)",
        "กฎ Tidy Data ที่ฝ่าฝืน",
        "ผลกระทบหลักต่อธุรกิจ & การสั่งการ AI"
    ]
    
    start_row = 6
    for col_idx, h in enumerate(headers, start=2):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[start_row].height = 28
    
    traps_meta = [
        (
            "กับดักที่ 1",
            "การบันทึกหลายค่ารวมในช่องเดียว\n(Multiple Values per Cell)",
            "ร้านปังปอนด์ โฮมเบเกอรี่",
            "01_ปังปอนด์เบเกอรี่ (หลายค่า)",
            "01_คำอธิบาย Trap 1",
            "กฎข้อที่ 3: 1 ช่อง = 1 ค่าเดียว\n(Each cell is a single value)",
            "สูตร =SUM() คำนวณไม่ได้เพราะมีตัวหนังสือปน / นับจำนวนชิ้นขนมแต่ละชนิดไม่ได้ / AI แยกราคาไม่ออก"
        ),
        (
            "กับดักที่ 2",
            "ยอดรวมแทรกกลางแถวข้อมูล\n(Subtotals Mixed in Rows)",
            "ศูนย์บริการ ProClean Car Care",
            "02_ProClean (ยอดรวมแทรก)",
            "02_คำอธิบาย Trap 2",
            "กฎข้อที่ 1: 1 แถว = 1 สิ่งที่สังเกต\n(Each row is an observation)",
            "สูตร =SUM() ทั้งคอลัมน์คำนวณเบิ้ล 2 เท่า / Pivot Table และ Looker Studio พังทันที / เรียงลำดับ Sort ไม่ได้"
        ),
        (
            "กับดักที่ 3",
            "เซลล์ผสานและหัวตารางซ้อน\n(Merged Cells & Multi-level Headers)",
            "บริษัท สยามฟู้ดส์ คอร์ปอเรชั่น (B2B)",
            "03_สยามฟู้ดส์ (เซลล์ผสาน)",
            "03_คำอธิบาย Trap 3",
            "กฎข้อที่ 2: 1 คอลัมน์ = 1 ตัวแปร\nและหัวตารางต้องมีเพียง 1 แถว",
            "เซลล์ล่างที่ถูกผสานกลายเป็นค่าว่าง (Null) ทันที / ตัวกรอง Filter หลุดหาย / นำเข้า Python หรือระบบ BI ไม่ได้"
        ),
        (
            "กับดักที่ 4",
            "การใช้ \"สี\" แทนข้อมูล\n(Color as Data)",
            "ร้านรองเท้าแฟชั่น Keng Footwear",
            "04_KengFootwear (สีแทนข้อมูล)",
            "04_คำอธิบาย Trap 4",
            "ต้องมีคอลัมน์สถานะ (Text Column)\nแทนการใช้สีไฮไลต์เซลล์",
            "สูตร Excel และ AI มองไม่เห็นสีเซลล์ / คำนวณยอดเงินตามสถานะไม่ได้ / โหลดเป็น CSV สีจะหายไป 100%"
        ),
        (
            "กับดักที่ 5",
            "ไม่มีรหัสกำกับเฉพาะ\n(Missing Unique Identifier)",
            "คลินิกทันตกรรม SmileCare",
            "05_SmileCare (ขาด UniqueID)",
            "05_คำอธิบาย Trap 5",
            "แต่ละแถวต้องมีรหัสอ้างอิงเฉพาะ\n(Primary Key / Unique ID)",
            "ลูกค้าชื่อ-นามสกุลซ้ำกัน ระบบเข้าใจผิดว่าเป็นคนเดียวกัน รวมยอดผิด / AI วิเคราะห์พฤติกรรมลูกค้าเพี้ยน"
        ),
        (
            "กับดักที่ 6",
            "ตารางแนวนอนที่ขยายตัวแปรออกขวา\n(Wide Format Trap)",
            "ร้านชาบูหม้อไฟ สุขสันต์ (5 สาขา)",
            "06_ชาบูสุขสันต์ (Wide Format)",
            "06_คำอธิบาย Trap 6",
            "ตารางหลังบ้านต้องเป็น Long Format\n(คอลัมน์คงที่ ขยายลงล่าง)",
            "ชื่อเดือนกลายเป็นหัวคอลัมน์ / เพิ่มเดือนใหม่ต้องแทรกคอลัมน์ทำให้สูตรพัง / สร้างกราฟแนวโน้มและ Pivot Table ลำบาก"
        ),
        (
            "กับดักที่ 7",
            "ข้อมูลไม่สม่ำเสมอและขาด Validation\n(Inconsistent & Unvalidated Formats)",
            "แบบฟอร์มสินค้าเกษตรแปรรูปชุมชน",
            "07_สินค้าเกษตร (ข้อมูลเพี้ยน)",
            "07_คำอธิบาย Trap 7",
            "การควบคุมคุณภาพข้อมูลต้นทาง\n(Quality at the Source / Validation)",
            "ชื่อจังหวัดแตกเป็น 5 รูปแบบ (กทม./Bangkok/BKK) / วันที่สะกดไม่เหมือนกันจัดเรียงไม่ได้ / ลืมกรอกเบอร์โทร Missing Data"
        )
    ]
    
    current_row = start_row + 1
    for item in traps_meta:
        ws.cell(row=current_row, column=2, value=item[0]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=2).font = bold_cell_font
        
        ws.cell(row=current_row, column=3, value=item[1]).font = bold_cell_font
        ws.cell(row=current_row, column=3).alignment = Alignment(vertical="center", wrap_text=True)
        
        ws.cell(row=current_row, column=4, value=item[2]).font = cell_font
        ws.cell(row=current_row, column=4).alignment = Alignment(vertical="center")
        
        c_data = ws.cell(row=current_row, column=5, value=item[3])
        c_data.font = tag_font
        c_data.fill = tag_fill
        c_data.alignment = Alignment(horizontal="center", vertical="center")
        
        c_exp = ws.cell(row=current_row, column=6, value=item[4])
        c_exp.font = tag_font
        c_exp.fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
        c_exp.alignment = Alignment(horizontal="center", vertical="center")
        
        ws.cell(row=current_row, column=7, value=item[5]).font = cell_font
        ws.cell(row=current_row, column=7).alignment = Alignment(vertical="center", wrap_text=True)
        
        ws.cell(row=current_row, column=8, value=item[6]).font = cell_font
        ws.cell(row=current_row, column=8).alignment = Alignment(vertical="center", wrap_text=True)
        
        # Apply borders and row height
        ws.row_dimensions[current_row].height = 42
        for col_idx in range(2, 9):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            if current_row % 2 == 1 and col_idx not in (5, 6):
                cell.fill = alt_fill
                
        current_row += 1
        
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 34
    ws.column_dimensions['D'].width = 28
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 25
    ws.column_dimensions['G'].width = 34
    ws.column_dimensions['H'].width = 46

def build_combined_workbook():
    master_wb = openpyxl.Workbook()
    
    # Sheet 1: Index Sheet
    index_ws = master_wb.active
    index_ws.title = "📋 สารบัญ & ภาพรวม 7 กับดัก"
    style_index_sheet(index_ws)
    
    # 7 File paths
    files_info = [
        ("bad_data_01_multiple_values_per_cell.xlsx", "01_ปังปอนด์เบเกอรี่ (หลายค่า)", "01_คำอธิบาย Trap 1"),
        ("bad_data_02_subtotals_in_rows.xlsx", "02_ProClean (ยอดรวมแทรก)", "02_คำอธิบาย Trap 2"),
        ("bad_data_03_merged_cells.xlsx", "03_สยามฟู้ดส์ (เซลล์ผสาน)", "03_คำอธิบาย Trap 3"),
        ("bad_data_04_color_as_data.xlsx", "04_KengFootwear (สีแทนข้อมูล)", "04_คำอธิบาย Trap 4"),
        ("bad_data_05_missing_unique_id.xlsx", "05_SmileCare (ขาด UniqueID)", "05_คำอธิบาย Trap 5"),
        ("bad_data_06_wide_format_trap.xlsx", "06_ชาบูสุขสันต์ (Wide Format)", "06_คำอธิบาย Trap 6"),
        ("bad_data_07_inconsistent_and_unvalidated.xlsx", "07_สินค้าเกษตร (ข้อมูลเพี้ยน)", "07_คำอธิบาย Trap 7"),
    ]
    
    for filename, new_data_title, new_exp_title in files_info:
        filepath = os.path.join(OUTPUT_DIR, filename)
        if not os.path.exists(filepath):
            print(f"File not found: {filepath}")
            continue
            
        src_wb = openpyxl.load_workbook(filepath, data_only=False)
        
        # In each source workbook, sheet[0] is data, sheet[1] is explanation
        data_ws_src = src_wb.worksheets[0]
        exp_ws_src = src_wb.worksheets[1] if len(src_wb.worksheets) > 1 else None
        
        # Clone data sheet into master_wb
        data_ws_dest = master_wb.create_sheet(title=new_data_title)
        copy_worksheet(data_ws_src, data_ws_dest)
        
        # Clone explanation sheet into master_wb
        if exp_ws_src:
            exp_ws_dest = master_wb.create_sheet(title=new_exp_title)
            copy_worksheet(exp_ws_src, exp_ws_dest)
            
    # Save master file
    master_path1 = os.path.join(OUTPUT_DIR, "bad_data_all_in_one.xlsx")
    master_path2 = os.path.join(OUTPUT_DIR, "bad_data_all_traps_collection.xlsx")
    
    master_wb.save(master_path1)
    master_wb.save(master_path2)
    print(f"Successfully created: {master_path1}")
    print(f"Successfully created: {master_path2}")

def copy_worksheet(source, target):
    target.views.sheetView[0].showGridLines = True
    
    # Copy column dimensions
    for col_letter, col_dim in source.column_dimensions.items():
        target.column_dimensions[col_letter].width = col_dim.width
        target.column_dimensions[col_letter].hidden = col_dim.hidden
        
    # Copy row dimensions
    for row_idx, row_dim in source.row_dimensions.items():
        target.row_dimensions[row_idx].height = row_dim.height
        target.row_dimensions[row_idx].hidden = row_dim.hidden
        
    # Copy cells with value, style, font, border, fill, alignment, number_format
    for row in source.iter_rows():
        for cell in row:
            new_cell = target.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font = Font(
                    name=cell.font.name,
                    size=cell.font.size,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    color=cell.font.color
                )
                new_cell.border = Border(
                    left=cell.border.left,
                    right=cell.border.right,
                    top=cell.border.top,
                    bottom=cell.border.bottom
                )
                if cell.fill and cell.fill.fill_type:
                    new_cell.fill = PatternFill(
                        fill_type=cell.fill.fill_type,
                        start_color=cell.fill.start_color,
                        end_color=cell.fill.end_color
                    )
                if cell.alignment:
                    new_cell.alignment = Alignment(
                        horizontal=cell.alignment.horizontal,
                        vertical=cell.alignment.vertical,
                        wrap_text=cell.alignment.wrap_text
                    )
                new_cell.number_format = cell.number_format
                
    # Copy merged cells
    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))

if __name__ == "__main__":
    build_combined_workbook()
