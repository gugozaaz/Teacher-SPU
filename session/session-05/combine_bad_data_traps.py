import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = r"d:\GoogleDrive\VibeCoding\teacher\session\session-05"

def style_index_sheet(ws):
    ws.views.sheetView[0].showGridLines = True
    
    title_font = Font(name="Noto Sans Thai", size=15, bold=True, color="1E1B4B")
    sub_font = Font(name="Noto Sans Thai", size=10, italic=True, color="4338CA")
    header_font = Font(name="Noto Sans Thai", size=10, bold=True, color="FFFFFF")
    cell_font = Font(name="Noto Sans Thai", size=10, color="1E293B")
    bold_cell_font = Font(name="Noto Sans Thai", size=10, bold=True, color="0F172A")
    tag_font = Font(name="Noto Sans Thai", size=9, bold=True, color="1D4ED8")
    tag_exp_font = Font(name="Noto Sans Thai", size=9, bold=True, color="15803D")
    
    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    tag_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    exp_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    ws["B2"] = "📚 รวมชุดข้อมูลตัวอย่าง 5 กับดัก Dirty Data (Session 05 Master Collection)"
    ws["B2"].font = title_font
    
    ws["B3"] = "รายวิชา: Fundamentals of Data Science | คณะการสร้างเจ้าของธุรกิจ มหาวิทยาลัยศรีปทุม"
    ws["B3"].font = sub_font
    
    ws["B4"] = "ชุดข้อมูลจำลองปัญหาข้อมูลสกปรกในธุรกิจไทย 5 ประเภท เพื่อฝึกฝนกระบวนการ Data Cleaning และเตรียมความพร้อมก่อนป้อนให้ AI"
    ws["B4"].font = Font(name="Noto Sans Thai", size=9.5, color="475569")
    
    headers = [
        "ลำดับ",
        "กับดักข้อผิดพลาด (Trap Name)",
        "กรณีศึกษาธุรกิจไทย",
        "แผ่นงานข้อมูลดิบ (Data Sheet)",
        "แผ่นงานคำอธิบาย (Explanation Sheet)",
        "กฎการคลีน / เครื่องมือหลัก",
        "ผลกระทบหลักต่อธุรกิจ & การสั่งการ AI"
    ]
    
    start_row = 6
    for col_idx, h in enumerate(headers, start=2):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[start_row].height = 28
    
    traps_meta = [
        (
            "กับดักที่ 1",
            "ข้อมูลซ้ำซ้อนจากหลายระบบ\n(Duplicate Records)",
            "ร้านชบา โคลทติ้ง (Chaba Clothing)",
            "01_ชบาโคลทติ้ง (ออเดอร์ซ้ำ)",
            "01_คำอธิบาย Trap 1",
            "กฎ Don't Touch Raw Data\nData > Remove duplicates หรือ =UNIQUE()",
            "ออเดอร์เบิ้ลจากการกดซ้ำหรือดึงไฟล์ซ้อน / ยอดขายและสต็อกพองเกินจริงกว่า 4,500 บาท / AI และ Pivot Table วิเคราะห์เบิ้ล 2 เท่า"
        ),
        (
            "กับดักที่ 2",
            "ช่องว่างแฝง & ตัวสะกดเพี้ยน\n(Hidden Whitespace & Typos)",
            "ร้านอาหารคลีน HealthyBites",
            "02_HealthyBites (ช่องว่างแฝง)",
            "02_คำอธิบาย Trap 2",
            "Data > Trim whitespace\nสูตร =TRIM(), =UPPER(), Find & Replace",
            "มี Space ข้างหน้า-หลังคำ / Pivot Table แยกเป็นคนละเมนู / ไซส์ตัวเล็กตัวใหญ่ปนกัน / AI รวมกลุ่มจัดหมวดหมู่ยอดขายไม่ได้"
        ),
        (
            "กับดักที่ 3",
            "ข้อมูลขาดหาย & เซลล์ว่างเปล่า\n(Missing Values & Nulls)",
            "ศูนย์ฟิตเนส GymFit Express",
            "03_GymFit (เซลล์ว่าง&Null)",
            "03_คำอธิบาย Trap 3",
            "Filter หาสาเหตุ / Conditional Formatting\nสูตร =ISBLANK() และเทคนิค Imputation",
            "เบอร์โทรศัพท์เว้นว่างติดต่อลูกค้าไม่ได้ / ช่องอายุคีย์ 'NULL' ปนตัวเลข / ค่าบริการเว้นว่างทำให้สูตรคำนวณยอดเงินรวมขาดหาย"
        ),
        (
            "กับดักที่ 4",
            "วันที่และตัวเลขฟอร์แมตมั่ว\n(Date & Number Format Chaos)",
            "ร้านแฟชั่นมัลติแบรนด์ (Retail)",
            "04_แฟชั่นรีเทล (ฟอร์แมตมั่ว)",
            "04_คำอธิบาย Trap 4",
            "Format > Number / Date\nสูตร =VALUE(), =DATEVALUE()",
            "ตัวเลขมีสัญลักษณ์ ฿, .- กลายเป็น Text ชิดซ้าย SUM ข้าม / วันที่ปี พ.ศ. ปน ค.ศ. เรียง Timeline ไม่ได้ / เบอร์โทร 0 หน้าหาย"
        ),
        (
            "กับดักที่ 5",
            "ค่าผิดปกติหลุดโลก & บิดเบือน\n(Outliers & Anomalies)",
            "ระบบคำสั่งซื้อ ชบา โคลทติ้ง",
            "05_ชบาออดิท (ค่าหลุดโลก)",
            "05_คำอธิบาย Trap 5",
            "Data Validation ป้องกันต้นทาง\nSort ตรวจหัว-ท้าย และ Conditional Formatting",
            "คีย์สั่งซื้อ 999 ชิ้น หรือราคาติดลบ / ส่วนลด 150% / บิดเบือนค่าเฉลี่ยและตัวเลขวิเคราะห์ของธุรกิจจน AI สรุปผิดทาง"
        )
    ]
    
    current_row = start_row + 1
    for item in traps_meta:
        ws.cell(row=current_row, column=2, value=item[0]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=2).font = bold_cell_font
        
        ws.cell(row=current_row, column=3, value=item[1]).font = bold_cell_font
        ws.cell(row=current_row, column=3).alignment = Alignment(vertical="center", wrap_text=True)
        
        ws.cell(row=current_row, column=4, value=item[2]).font = cell_font
        ws.cell(row=current_row, column=4).alignment = Alignment(vertical="center")
        
        c_data = ws.cell(row=current_row, column=5, value=item[3])
        c_data.font = tag_font
        c_data.fill = tag_fill
        c_data.alignment = Alignment(horizontal="center", vertical="center")
        c_data.hyperlink = f"#'{item[3]}'!A1"
        
        c_exp = ws.cell(row=current_row, column=6, value=item[4])
        c_exp.font = tag_exp_font
        c_exp.fill = exp_fill
        c_exp.alignment = Alignment(horizontal="center", vertical="center")
        c_exp.hyperlink = f"#'{item[4]}'!A1"
        
        ws.cell(row=current_row, column=7, value=item[5]).font = cell_font
        ws.cell(row=current_row, column=7).alignment = Alignment(vertical="center", wrap_text=True)
        
        ws.cell(row=current_row, column=8, value=item[6]).font = cell_font
        ws.cell(row=current_row, column=8).alignment = Alignment(vertical="center", wrap_text=True)
        
        # Apply borders and row height
        ws.row_dimensions[current_row].height = 42
        for col_idx in range(2, 9):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            if current_row % 2 == 1 and col_idx not in (5, 6):
                cell.fill = alt_fill
                
        current_row += 1
        
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 34
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 32
    ws.column_dimensions['F'].width = 25
    ws.column_dimensions['G'].width = 36
    ws.column_dimensions['H'].width = 46

def copy_worksheet(source, target):
    target.views.sheetView[0].showGridLines = True
    
    # Copy column dimensions
    for col_letter, col_dim in source.column_dimensions.items():
        target.column_dimensions[col_letter].width = col_dim.width
        target.column_dimensions[col_letter].hidden = col_dim.hidden
        
    # Copy row dimensions
    for row_idx, row_dim in source.row_dimensions.items():
        target.row_dimensions[row_idx].height = row_dim.height
        target.row_dimensions[row_idx].hidden = row_dim.hidden
        
    # Copy cells with value, style, font, border, fill, alignment, number_format
    for row in source.iter_rows():
        for cell in row:
            new_cell = target.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font = Font(
                    name=cell.font.name,
                    size=cell.font.size,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    color=cell.font.color
                )
                new_cell.border = Border(
                    left=cell.border.left,
                    right=cell.border.right,
                    top=cell.border.top,
                    bottom=cell.border.bottom
                )
                if cell.fill and cell.fill.fill_type:
                    new_cell.fill = PatternFill(
                        fill_type=cell.fill.fill_type,
                        start_color=cell.fill.start_color,
                        end_color=cell.fill.end_color
                    )
                if cell.alignment:
                    new_cell.alignment = Alignment(
                        horizontal=cell.alignment.horizontal,
                        vertical=cell.alignment.vertical,
                        wrap_text=cell.alignment.wrap_text
                    )
                new_cell.number_format = cell.number_format
                
    # Copy merged cells
    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))

def build_combined_workbook():
    master_wb = openpyxl.Workbook()
    
    # Sheet 1: Index Sheet
    index_ws = master_wb.active
    index_ws.title = "📋 สารบัญ & ภาพรวม 5 กับดัก"
    style_index_sheet(index_ws)
    
    # 5 Trap files matching Session 4 structure & <= 31 chars length
    files_info = [
        ("bad_data_01_duplicate_records.xlsx", "01_ชบาโคลทติ้ง (ออเดอร์ซ้ำ)", "01_คำอธิบาย Trap 1"),
        ("bad_data_02_hidden_whitespace_and_typos.xlsx", "02_HealthyBites (ช่องว่างแฝง)", "02_คำอธิบาย Trap 2"),
        ("bad_data_03_missing_values_and_nulls.xlsx", "03_GymFit (เซลล์ว่าง&Null)", "03_คำอธิบาย Trap 3"),
        ("bad_data_04_date_and_number_format_chaos.xlsx", "04_แฟชั่นรีเทล (ฟอร์แมตมั่ว)", "04_คำอธิบาย Trap 4"),
        ("bad_data_05_outliers_and_anomalies.xlsx", "05_ชบาออดิท (ค่าหลุดโลก)", "05_คำอธิบาย Trap 5"),
    ]
    
    for filename, new_data_title, new_exp_title in files_info:
        filepath = os.path.join(OUTPUT_DIR, filename)
        if not os.path.exists(filepath):
            print(f"File not found: {filepath}")
            continue
            
        src_wb = openpyxl.load_workbook(filepath, data_only=False)
        
        # In each source workbook, sheet[0] is data, sheet[1] is explanation
        data_ws_src = src_wb.worksheets[0]
        exp_ws_src = src_wb.worksheets[1] if len(src_wb.worksheets) > 1 else None
        
        # Clone data sheet into master_wb
        data_ws_dest = master_wb.create_sheet(title=new_data_title)
        copy_worksheet(data_ws_src, data_ws_dest)
        
        # Clone explanation sheet into master_wb
        if exp_ws_src:
            exp_ws_dest = master_wb.create_sheet(title=new_exp_title)
            copy_worksheet(exp_ws_src, exp_ws_dest)
            
    # Save target files matching naming standard
    out_path1 = os.path.join(OUTPUT_DIR, "bad_data_01_to_05_combined.xlsx")
    out_path2 = os.path.join(OUTPUT_DIR, "bad_data_all_in_one.xlsx")
    out_path3 = os.path.join(OUTPUT_DIR, "bad_data_5_traps_combined.xlsx")
    
    master_wb.save(out_path1)
    master_wb.save(out_path2)
    master_wb.save(out_path3)
    
    print(f"Successfully created: {out_path1}")
    print(f"Successfully created: {out_path2}")
    print(f"Successfully created: {out_path3}")

if __name__ == "__main__":
    build_combined_workbook()
