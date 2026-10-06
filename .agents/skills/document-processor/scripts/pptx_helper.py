#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
PowerPoint (.pptx) Helper Script for Antigravity
Provides utilities for creating modern presentations, modifying slides, and manipulating shapes.
"""

import sys
import os
import io
import argparse
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Force UTF-8
if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')


def create_modern_presentation(output_path: str, title: str, subtitle: str = None, theme: str = "dark"):
    prs = Presentation()
    # 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Theme colors
    if theme == "dark":
        bg_rgb = RGBColor(0x0F, 0x17, 0x2A)       # Deep slate
        title_rgb = RGBColor(0x38, 0xBD, 0xF8)    # Cyan / Sky
        text_rgb = RGBColor(0xF8, 0xFA, 0xFC)     # White slate
        sub_rgb = RGBColor(0x94, 0xA3, 0xB8)      # Muted slate
    else:
        bg_rgb = RGBColor(0xF8, 0xFA, 0xFC)       # Clean light
        title_rgb = RGBColor(0x02, 0x84, 0xC7)    # Blue
        text_rgb = RGBColor(0x0F, 0x17, 0x2A)     # Dark slate
        sub_rgb = RGBColor(0x64, 0x74, 0x8B)      # Slate
        
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # Background shape
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_rgb
    bg.line.fill.background()
    
    # Title & Subtitle in single text frame
    tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.0))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = 'Segoe UI'
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = title_rgb
    p.alignment = PP_ALIGN.LEFT
    
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = 'Segoe UI'
        p2.font.size = Pt(22)
        p2.font.color.rgb = sub_rgb
        p2.space_before = Pt(16)
        p2.alignment = PP_ALIGN.LEFT
        
    prs.save(output_path)
    print(f"Created presentation ({theme} theme): {output_path}")
    return prs


def add_content_slide(prs, title: str, bullet_points: list, theme: str = "dark"):
    if theme == "dark":
        bg_rgb = RGBColor(0x0F, 0x17, 0x2A)
        title_rgb = RGBColor(0x38, 0xBD, 0xF8)
        text_rgb = RGBColor(0xF8, 0xFA, 0xFC)
    else:
        bg_rgb = RGBColor(0xF8, 0xFA, 0xFC)
        title_rgb = RGBColor(0x02, 0x84, 0xC7)
        text_rgb = RGBColor(0x0F, 0x17, 0x2A)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_rgb
    bg.line.fill.background()
    
    # Slide Title
    t_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.0))
    t_frame = t_box.text_frame
    tp = t_frame.paragraphs[0]
    tp.text = title
    tp.font.name = 'Segoe UI'
    tp.font.size = Pt(32)
    tp.font.bold = True
    tp.font.color.rgb = title_rgb
    
    # Body Bullets
    b_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.8))
    b_frame = b_box.text_frame
    b_frame.word_wrap = True
    
    for i, pt_text in enumerate(bullet_points):
        p = b_frame.paragraphs[0] if i == 0 else b_frame.add_paragraph()
        p.text = f"•  {pt_text}"
        p.font.name = 'Segoe UI'
        p.font.size = Pt(18)
        p.font.color.rgb = text_rgb
        p.space_after = Pt(14)
        
    return slide


def main():
    parser = argparse.ArgumentParser(description="PowerPoint (.pptx) Helper")
    subparsers = parser.add_subparsers(dest="action", required=True)
    
    # Create presentation
    create_p = subparsers.add_parser("create", help="Create a new deck")
    create_p.add_argument("output", help="Output .pptx path")
    create_p.add_argument("--title", required=True, help="Title text")
    create_p.add_argument("--subtitle", help="Subtitle text")
    create_p.add_argument("--theme", choices=["dark", "light"], default="dark", help="Theme")
    
    args = parser.parse_args()
    
    if args.action == "create":
        create_modern_presentation(args.output, args.title, args.subtitle, args.theme)


if __name__ == "__main__":
    main()
