#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
PDF to PowerPoint (.pptx) Converter for Antigravity
Converts PDF slides into high-fidelity 16:9 widescreen PowerPoint presentations with extracted text in notes.
"""

import sys
import os
import io
import argparse
import tempfile
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt

# Force UTF-8
if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')


def convert_pdf_to_pptx(pdf_path: str, pptx_path: str, dpi: int = 200):
    pdf_path = os.path.abspath(pdf_path)
    pptx_path = os.path.abspath(pptx_path)
    
    if not os.path.exists(pdf_path):
        print(f"Error: PDF file not found: {pdf_path}", file=sys.stderr)
        sys.exit(1)
        
    doc = pymupdf.open(pdf_path)
    total_pages = len(doc)
    print(f"Reading PDF with {total_pages} page(s)...")
    
    # Create 16:9 PowerPoint Presentation (13.333" x 7.5")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]
    
    with tempfile.TemporaryDirectory() as tmp_dir:
        for idx, page in enumerate(doc, start=1):
            # Render page to high-res image
            zoom = dpi / 72.0
            mat = pymupdf.Matrix(zoom, zoom)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            
            img_path = os.path.join(tmp_dir, f"slide_{idx:03d}.png")
            pix.save(img_path)
            
            # Add slide to presentation
            slide = prs.slides.add_slide(blank_slide_layout)
            
            # Insert full-bleed background image
            slide.shapes.add_picture(
                img_path,
                Inches(0),
                Inches(0),
                width=Inches(13.333),
                height=Inches(7.5)
            )
            
            # Extract text from page and add to speaker notes for searchability & accessibility
            page_text = page.get_text().strip()
            if page_text and slide.notes_slide:
                tf = slide.notes_slide.notes_text_frame
                tf.text = f"[Slide {idx} Content]\n{page_text}"
                
            print(f"Processed Slide {idx}/{total_pages}...")
            
    prs.save(pptx_path)
    print(f"\nSuccessfully generated PowerPoint presentation: {pptx_path}")


def main():
    parser = argparse.ArgumentParser(description="Convert PDF to PowerPoint Presentation (.pptx)")
    parser.add_argument("pdf_file", help="Input PDF file path")
    parser.add_argument("pptx_file", help="Output .pptx file path")
    parser.add_argument("--dpi", type=int, default=200, help="Rendering DPI (default: 200 for crisp 1080p+ quality)")
    
    args = parser.parse_args()
    convert_pdf_to_pptx(args.pdf_file, args.pptx_file, args.dpi)


if __name__ == "__main__":
    main()
